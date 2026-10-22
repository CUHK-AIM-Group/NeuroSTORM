#!/usr/bin/env python
"""Draw Fig.5 as Source Data trial points or measured query-scan violins.

Points mode reads the Top1 and map sheets of a supplied Source Data workbook.
Violin mode reads the eight ReID ranking lists. Its dots and box statistics
use measured query-scan scores; KDE smooths these scores for display.
"""

from __future__ import annotations

import argparse
import json
import os
from numbers import Real
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.patches import Rectangle
import numpy as np


BASE = Path(__file__).resolve().parents[1]
DEFAULT_PROTOCOL = BASE / "configs/reid/fixed400.json"
MODELS = ("NeuroSTORM", "SwiFT", "TFF", "BrainLM", "Com-brainTF", "LG-GNN", "BNT", "BrainGNN")
SCAN_NUMBERS = {"rfMRI_REST1_LR": 1, "rfMRI_REST1_RL": 2,
                "rfMRI_REST2_LR": 3, "rfMRI_REST2_RL": 4}
COLORS = {
    "top1": {
        "BrainGNN": [0.204, 0.294, 0.227], "BNT": [0.208, 0.388, 0.259],
        "LG-GNN": [0.212, 0.486, 0.290], "Com-brainTF": [0.216, 0.580, 0.318],
        "BrainLM": [0.267, 0.655, 0.369], "TFF": [0.373, 0.710, 0.439],
        "SwiFT": [0.478, 0.765, 0.514], "NeuroSTORM": [0.580, 0.820, 0.584],
    },
    "mAP": {
        "BrainGNN": [0.204, 0.271, 0.322], "BNT": [0.208, 0.345, 0.443],
        "LG-GNN": [0.212, 0.420, 0.569], "Com-brainTF": [0.216, 0.490, 0.690],
        "BrainLM": [0.267, 0.561, 0.773], "TFF": [0.369, 0.624, 0.804],
        "SwiFT": [0.471, 0.690, 0.835], "NeuroSTORM": [0.573, 0.753, 0.867],
    },
}
POINT_COLORS = ("#B64342", "#315E91", "#62809A", "#796391",
                "#7A9788", "#B38B58", "#666666", "#6A9B9E")
POINT_MARKERS = ("o", "s", "^", "D", "v", "P", "X", ">")
POINT_STYLES = ("-", "-", "--", "-.", (0, (4, 2)),
                (0, (2, 1)), ":", (0, (6, 2, 1, 2)))


def load_source_data(source_path):
    """Read the supplied five rows verbatim; each cell is one plotted value."""
    import openpyxl

    panels = []
    book = openpyxl.load_workbook(source_path, read_only=True, data_only=True)
    try:
        for sheet_name, metric in (("Top1", "top1"), ("map", "mAP")):
            if sheet_name not in book.sheetnames:
                raise ValueError(f"Source Data is missing sheet {sheet_name!r}")
            sheet = book[sheet_name]
            if (sheet.max_row, sheet.max_column) != (6, 9):
                raise ValueError(f"{sheet_name}: expected five trials and eight model columns")
            models = tuple(sheet.cell(1, c).value for c in range(2, 10))
            if len(set(models)) != len(MODELS) or set(models) != set(MODELS):
                raise ValueError(f"{sheet_name}: expected the eight ReID model names")
            values = np.empty((5, len(MODELS)), dtype=float)
            for trial in range(5):
                for index, model in enumerate(MODELS):
                    cell = sheet.cell(trial + 2, models.index(model) + 2)
                    number = cell.value
                    if (not isinstance(number, Real) or isinstance(number, bool)
                            or not np.isfinite(number) or not 0 <= number <= 100):
                        raise ValueError(f"{sheet_name}!{cell.coordinate}: expected a percentage from 0 to 100")
                    values[trial, index] = number
            panels.append(dict(sheet=sheet_name, metric=metric, models=MODELS, values=values))
    finally:
        book.close()
    return panels


def scan_identity(scan_id):
    subject, sequence = scan_id.split("::")
    return int(subject), SCAN_NUMBERS[sequence]


def load_ranking_lists(ranking_dir, protocol_path=DEFAULT_PROTOCOL):
    protocol = json.loads(Path(protocol_path).read_text(encoding="utf-8-sig"))
    allowed = set(protocol["fixed400_query_scans"])
    rows = []
    for model in MODELS:
        occurrences, trial_means = {}, []
        path = Path(ranking_dir) / (model + ".txt")
        with path.open(encoding="utf-8-sig") as handle:
            for trial in protocol["test"]:
                queries = [scan for scan in trial["query_scans"] if scan in allowed]
                gallery_ids = {scan_identity(scan) for scan in trial["gallery_scans"]}
                scores = []
                for query in queries:
                    subject, number = scan_identity(query)
                    if handle.readline().split() != [str(subject), str(number)]:
                        raise ValueError(f"{model}: query order does not match the saved protocol")
                    ranked = json.loads(handle.readline())
                    if len(ranked) != len(gallery_ids) or any(len(candidate) != 4 for candidate in ranked):
                        raise ValueError(f"{model}: expected complete four-field gallery rankings")
                    if {(candidate[0], candidate[1]) for candidate in ranked} != gallery_ids:
                        raise ValueError(f"{model}: gallery scans do not match the saved protocol")
                    similarities = np.asarray([candidate[2] for candidate in ranked], dtype=np.float64)
                    if not np.isfinite(similarities).all() or np.any(np.diff(similarities) > 0):
                        raise ValueError(f"{model}: similarities must be finite and in descending order")
                    relevant = np.asarray([candidate[0] == subject for candidate in ranked])
                    if any(candidate[3] != int(hit) for candidate, hit in zip(ranked, relevant)):
                        raise ValueError(f"{model}: relevance flags do not match subject identities")
                    hit_ranks = np.flatnonzero(relevant) + 1
                    if protocol.get("allow_single_positive_gallery") is True:
                        if len(hit_ranks) not in (1, 2):
                            raise ValueError(f"{model}: each query must have one or two relevant gallery scans")
                    elif len(hit_ranks) != 2:
                        raise ValueError(f"{model}: each query must have two relevant gallery scans")
                    top1 = float(relevant[0])
                    ap = float(np.mean(np.arange(1, len(hit_ranks) + 1) / hit_ranks))
                    scores.append((top1, ap))
                    occurrences.setdefault(query, []).append((top1, ap))
                trial_means.append(np.mean(scores, axis=0))
            if handle.read().strip():
                raise ValueError(f"{model}: unexpected query records after the saved trials")
        if set(occurrences) != allowed:
            raise ValueError(f"{model}: rankings do not cover the fixed query scans")
        points = []
        for query in sorted(occurrences):
            top1, ap = np.mean(occurrences[query], axis=0) * 100
            points.append(dict(scan_id=query, top1_percent=float(top1), mAP_percent=float(ap)))
        top1, ap = np.mean(trial_means, axis=0)
        split_counts = protocol.get("subject_split", {})
        participant_counts = dict(validation=split_counts["val"], test=split_counts["test"]) if split_counts else None
        rows.append(dict(model=model, points=points, metrics=dict(top1=float(top1), mAP=float(ap)),
                         participant_counts=participant_counts))
    return rows


def vertical_jitter(scan_id, y):
    number = np.random.default_rng(list(scan_id.encode("utf-8"))).random()
    return y + 0.015 - number * 0.30


def query_density(values, bandwidth_pp=2.6):
    """Smooth measured scores for display, reflecting kernels at 0 and 100.

    The fixed bandwidth is in percentage points; it changes only the curve,
    not the measured dots or the box statistics.
    """
    grid = np.linspace(0, 100, 501)
    curve = np.zeros_like(grid)
    for centers in (values, -values, 200 - values):
        offsets = (grid[:, None] - centers[None, :]) / bandwidth_pp
        curve += np.exp(-0.5 * offsets ** 2).sum(axis=1)
    return grid, curve / curve.max()


def query_point_heights(values, point_ids, y):
    """Spread tied scores vertically without moving their measured x values."""
    heights = np.empty(len(values))
    for value in np.unique(values):
        indices = np.flatnonzero(values == value)
        ordered = sorted(indices, key=lambda index: point_ids[index])
        heights[ordered] = (vertical_jitter(point_ids[ordered[0]], y) if len(ordered) == 1
                            else y + np.linspace(-0.27, -0.01, len(ordered)))
    return heights


def box_and_whiskers(ax, values, y):
    q1, median, q3 = np.percentile(values, [25, 50, 75])
    iqr = q3 - q1
    inliers = (values >= q1 - 1.5 * iqr) & (values <= q3 + 1.5 * iqr)
    low, high = values[inliers].min(), values[inliers].max()
    ax.plot([low, high], [y, y], color="#333333", lw=0.30, zorder=3)
    ax.vlines([low, high], y - 0.025, y + 0.025, color="#333333", lw=0.30, zorder=3)
    ax.add_patch(Rectangle((q1, y - 0.04), q3 - q1, 0.08, facecolor="none",
                           edgecolor="#333333", linewidth=0.30, zorder=4))
    ax.vlines(median, y - 0.04, y + 0.04, color="#333333", lw=0.35, zorder=4)
    return inliers


def legend(fig, width, height, axes_top, query_size, participant_counts=None):
    left, right = (44.010 - 34.0) / width, (552.891 - 34.0) / width
    bottom = (axes_top + 6.363) / height
    rect_height = 11.122 / height
    fig.add_artist(Rectangle((left, bottom), right - left, rect_height,
                             transform=fig.transFigure, fill=False, linewidth=0.45, edgecolor="#303030"))
    count_suffix = f" = {query_size} query"
    if participant_counts is not None:
        validation, test = participant_counts["validation"], participant_counts["test"]
        count_suffix = f" = {validation+test} ({query_size} query)"
    groups = [
        [("Violin width:", "bold", "normal"), (" kernel density", "normal", "normal")],
        [("Inner box/line:", "bold", "normal"), (" median and interquartile range", "normal", "normal")],
        [("Whiskers:", "bold", "normal"), (" 1.5×IQR", "normal", "normal")],
        [("Dots:", "bold", "normal"), (" query scan", "normal", "normal")],
        [("Circles:", "bold", "normal"), (" outliers", "normal", "normal")],
        [("n", "normal", "italic"), (count_suffix, "normal", "normal")],
    ]
    artists = [[fig.text(0, bottom+.32*rect_height, text, fontsize=5.5,
                         fontfamily="Times New Roman", fontweight=weight, fontstyle=style,
                         va="baseline", ha="left", color="#161616")
                for text, weight, style in group] for group in groups]
    fig.canvas.draw()
    renderer = fig.canvas.get_renderer()
    widths = [[artist.get_window_extent(renderer).width/fig.bbox.width for artist in group]
              for group in artists]
    inner_left, inner_right = left + 5.5 / width, right - 5.5 / width
    gap = (inner_right-inner_left-sum(map(sum, widths)))/(len(groups)-1)
    assert gap >= 0, "Legend text exceeds the measured paper legend."
    x = inner_left
    for group, group_widths in zip(artists, widths):
        for artist, artist_width in zip(group, group_widths):
            artist.set_x(x)
            x += artist_width
        x += gap


def save_figure(fig, output_dir, stem):
    for extension in ("png", "svg", "pdf"):
        fig.savefig(output_dir / f"{stem}.{extension}", dpi=600 if extension == "png" else None,
                    facecolor="white", metadata={"Creator": "plot_fig5.py"} if extension == "pdf" else None)
    plt.close(fig)


def draw_points(panels, output_dir):
    """Plot the five Source Data rows with the supplied reference styling."""
    fig, axes = plt.subplots(1, 2, figsize=(183 / 25.4, 76 / 25.4), sharey=True)
    fig.subplots_adjust(left=0.082, right=0.98, bottom=0.18, top=0.86, wspace=0.26)
    x = np.arange(1, 6)
    series = []
    available_fonts = {font.name for font in font_manager.fontManager.ttflist}
    panel_font = next(name for name in ("Arial Black", "Arial", "DejaVu Sans")
                      if name in available_fonts)
    for panel_id, (ax, panel) in enumerate(zip(axes, panels)):
        lines = []
        for index, model in enumerate(MODELS):
            line, = ax.plot(
                x, panel["values"][:, index], label=model,
                color=POINT_COLORS[index], marker=POINT_MARKERS[index],
                linestyle=POINT_STYLES[index],
                linewidth=1.7 if index == 0 else 1.0,
                markersize=4.2 if index == 0 else 3.5,
                markerfacecolor=POINT_COLORS[index] if index == 0 else "white",
                markeredgewidth=0.9, zorder=5 if index == 0 else 3,
            )
            line.set_gid(f"source-series-{panel['metric']}-{model.lower().replace('-', '')}")
            lines.append(line)
        series.append(lines)
        ax.text(-0.17, 1.035, "ab"[panel_id], transform=ax.transAxes,
                fontfamily=panel_font, fontsize=9.3,
                fontweight="black", va="bottom")
        ax.set_xlim(0.8, 5.2)
        ax.set_ylim(75, 100)
        ax.set_xticks(x, [f"Trial {i}" for i in x])
        ax.set_yticks(np.arange(75, 101, 5))
        ax.set_xlabel("HCP-YA (n = 216)", labelpad=6, fontsize=8)
        label = "Top-1 accuracy (%)" if panel_id == 0 else "Mean average precision (%)"
        ax.set_ylabel(label, labelpad=7, fontsize=8)
        ax.grid(axis="y", color="#E5E5E5", linewidth=0.25, zorder=0)
        ax.tick_params(axis="both", length=2.7, width=0.25, pad=3, labelsize=7)
        ax.tick_params(axis="y", labelleft=True)
        for side in ("top", "right"):
            ax.spines[side].set_visible(False)
        for side in ("left", "bottom"):
            ax.spines[side].set_color("black")
            ax.spines[side].set_linewidth(0.25)
    fig.legend(series[0], MODELS, loc="upper center", bbox_to_anchor=(0.53, 0.98), ncol=8,
               columnspacing=1.0, handlelength=2.2, handletextpad=0.4,
               labelspacing=0.7, borderaxespad=0, fontsize=7, frameon=False)
    save_figure(fig, output_dir, "Fig5_points")


def draw_violin(rows, output_dir):
    display_order = ("BrainGNN", "BNT", "LG-GNN", "Com-brainTF", "BrainLM", "TFF", "SwiFT", "NeuroSTORM")
    rows = sorted(rows, key=lambda row: display_order.index(row["model"]))
    query_size = len(rows[0]["points"])
    if not all(len(row["points"]) == query_size for row in rows):
        raise ValueError("All models must represent the same query scans")
    n = len(rows)
    width = 530.0
    bottom = 25.316
    axes_height = 126.016 * n / 8
    axes_top = bottom + axes_height
    height = axes_top + 22.668
    fig = plt.figure(figsize=(width / 72, height / 72), dpi=150)
    left_rect = [(85.641 - 34.0) / width, bottom / height, 200.587 / width, axes_height / height]
    right_rect = [(345.289 - 34.0) / width, bottom / height, 200.974 / width, axes_height / height]
    for rect, metric, xlabel in [
        (left_rect, "top1", "Top-1 accuracy"),
        (right_rect, "mAP", "Mean average precision"),
    ]:
        ax = fig.add_axes(rect)
        for i, row in enumerate(rows):
            y = n - i - 1
            values = np.asarray([point[metric + "_percent"] for point in row["points"]])
            point_ids = [point["scan_id"] for point in row["points"]]
            grid, curve = query_density(values)
            ax.fill_between(grid, y + 0.065, y + 0.065 + 0.30 * curve,
                            facecolor=COLORS[metric][row["model"]], edgecolor="#555555",
                            linewidth=0.35, zorder=1)
            mask = box_and_whiskers(ax, values, y + 0.04)
            ys = query_point_heights(values, point_ids, y)
            model_id = row["model"].lower().replace("-", "")
            points = ax.scatter(values[mask], ys[mask], s=0.70, c="#555555",
                                alpha=0.42, linewidths=0, zorder=5)
            points.set_gid(f"scans-inliers-{metric}-{model_id}")
            if (~mask).any():
                rings = ax.scatter(values[~mask], ys[~mask], s=1.5, marker="o", facecolors="none",
                                   edgecolors="#424242", linewidths=0.30, zorder=6)
                rings.set_gid(f"scans-outliers-{metric}-{model_id}")
        ax.set(xlim=(-1.35, 101.35), ylim=(-0.5, n - 0.5), yticks=np.arange(n - 1, -1, -1),
               xticks=np.arange(0, 101, 20))
        ax.set_yticklabels([row["model"] for row in rows])
        for spine in ax.spines.values():
            spine.set_linewidth(0.35)
            spine.set_color("#606060")
        ax.tick_params(axis="both", length=2.6, width=0.35, color="#606060", pad=2.7)
        ax.tick_params(axis="x", labelsize=5.5)
        ax.tick_params(axis="y", labelsize=6.5)
        ax.grid(False)
        fig.text(rect[0] + rect[2] / 2, (262.928 - 257.0) / height, xlabel,
                 fontsize=6.5, ha="center", va="baseline")
    legend(fig, width, height, axes_top, query_size, rows[0].get("participant_counts"))
    save_figure(fig, output_dir, "Fig5_violin")


def draw(data, output_dir, mode="violin"):
    if mode == "points":
        with plt.rc_context({"font.family": "sans-serif",
                             "font.sans-serif": ["Arial", "DejaVu Sans", "Liberation Sans"],
                             "text.color": "black", "axes.labelcolor": "black",
                             "xtick.color": "black", "ytick.color": "black"}):
            draw_points(data, output_dir)
    elif mode == "violin":
        draw_violin(data, output_dir)
    else:
        raise ValueError(f"Unknown plotting mode: {mode}")


def register_paper_fonts(mode="violin"):
    system_fonts = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"
    names = (("arial.ttf", "arialbd.ttf", "ariali.ttf", "arialbi.ttf", "ariblk.ttf")
             if mode == "points" else ("times.ttf", "timesbd.ttf", "timesi.ttf", "timesbi.ttf"))
    for name in names:
        if (system_fonts/name).is_file():
            font_manager.fontManager.addfont(system_fonts/name)
    if mode == "violin":
        font_manager.findfont(font_manager.FontProperties(family="Times New Roman"), fallback_to_default=False)
        plt.rcParams["font.family"] = "Times New Roman"
    else:
        plt.rcParams.update({"font.family": "sans-serif",
                             "font.sans-serif": ["Arial", "DejaVu Sans", "Liberation Sans"]})
    plt.rcParams.update({"pdf.fonttype": 42, "ps.fonttype": 42,
                         "svg.fonttype": "path", "axes.unicode_minus": False, "text.color": "#161616"})


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--mode", choices=("points", "violin"), default="violin")
    parser.add_argument("--source", type=Path,
                        help="Source Data Excel workbook for points mode (Top1 and map sheets)")
    parser.add_argument("--ranking-dir", type=Path,
                        help="Local directory containing the eight ranking lists for violin mode")
    parser.add_argument("--protocol", type=Path, default=DEFAULT_PROTOCOL)
    parser.add_argument("--output-dir", type=Path, default=Path("results/fig5"))
    args = parser.parse_args()
    if args.mode == "points":
        if args.source is None:
            parser.error("--source is required for points mode")
        data = load_source_data(args.source)
    else:
        if args.ranking_dir is None:
            parser.error("--ranking-dir is required for violin mode")
        data = load_ranking_lists(args.ranking_dir, args.protocol)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    register_paper_fonts(args.mode)
    draw(data, args.output_dir, args.mode)
    if args.mode == "points":
        for index, model in enumerate(MODELS):
            top1 = data[0]["values"][:, index].mean()
            ap = data[1]["values"][:, index].mean()
            print(f"{model}: Source Data mean Top-1 {top1:.6f}%, mAP {ap:.6f}%")
    else:
        for row in data:
            metrics = row["metrics"]
            print(f"{row['model']}: Top-1 {metrics['top1'] * 100:.6f}%, mAP {metrics['mAP'] * 100:.6f}%")
    print(f"Figure saved in {args.output_dir}")


if __name__ == "__main__":
    main()
