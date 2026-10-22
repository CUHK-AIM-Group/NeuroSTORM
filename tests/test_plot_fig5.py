"""Check that Fig.5 represents measured query scores and trial statistics."""
import copy
import json
import sys

from matplotlib import rc_context
from matplotlib.figure import Figure
import numpy as np
import openpyxl
import pytest

from scripts.task4 import plot_fig5


@pytest.fixture
def ranking_files(tmp_path):
    queries = ["100001::rfMRI_REST1_LR", "100002::rfMRI_REST1_LR"]
    gallery = ["100001::rfMRI_REST2_LR", "100001::rfMRI_REST2_RL",
               "100002::rfMRI_REST2_LR", "100002::rfMRI_REST2_RL"]
    protocol = {"fixed400_query_scans": queries, "test": [
        {"query_scans": queries, "gallery_scans": gallery},
        {"query_scans": queries[:1], "gallery_scans": gallery},
    ]}
    protocol_path = tmp_path / "protocol.json"
    protocol_path.write_text(json.dumps(protocol), encoding="utf-8")
    records = []
    for query, order in [(queries[0], [0, 2, 3, 1]),
                         (queries[1], [0, 2, 3, 1]),
                         (queries[0], [0, 1, 2, 3])]:
        subject, number = plot_fig5.scan_identity(query)
        ranked = []
        for rank, index in enumerate(order):
            candidate, scan_number = plot_fig5.scan_identity(gallery[index])
            ranked.append([candidate, scan_number, .9 - .1 * rank, int(candidate == subject)])
        records.extend([f"{subject} {number}", json.dumps(ranked)])
    for model in plot_fig5.MODELS:
        (tmp_path / f"{model}.txt").write_text("\n".join(records) + "\n", encoding="utf-8")
    return tmp_path, protocol_path


def test_query_scores_and_trial_means_have_distinct_weights(ranking_files):
    rows = plot_fig5.load_ranking_lists(*ranking_files)
    for row in rows:
        points = row["points"]
        assert [point["top1_percent"] for point in points] == [100., 0.]
        assert [point["mAP_percent"] for point in points] == pytest.approx([87.5, 100 * 7 / 12])
        assert row["metrics"] == pytest.approx({"top1": .75, "mAP": 5 / 6})
        assert row["metrics"]["top1"] != np.mean([point["top1_percent"] for point in points]) / 100


def test_inconsistent_relevance_flags_are_rejected(ranking_files):
    directory, protocol = ranking_files
    path = directory / "NeuroSTORM.txt"
    lines = path.read_text(encoding="utf-8").splitlines()
    ranked = json.loads(lines[1])
    ranked[0][3] = 0
    lines[1] = json.dumps(ranked)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    with pytest.raises(ValueError, match="relevance flags"):
        plot_fig5.load_ranking_lists(directory, protocol)


def test_drawing_preserves_all_measured_scores_and_method_order(tmp_path, monkeypatch):
    rows = [dict(model=model, metrics=dict(top1=.8, mAP=.7), points=[
        dict(scan_id=f"{100001 + i}::rfMRI_REST1_LR", top1_percent=top1, mAP_percent=ap)
        for i, (top1, ap) in enumerate(zip([0, 100, 100, 100, 100], [0, 50, 100, 100, 100]))
    ]) for model in plot_fig5.MODELS]
    original = copy.deepcopy(rows)
    saved = []
    monkeypatch.setattr(Figure, "savefig", lambda fig, path, **kwargs: saved.append((fig, path)))
    # This score/order check uses Matplotlib's bundled font on every platform.
    with rc_context({"font.family": "DejaVu Serif"}):
        plot_fig5.draw(rows, tmp_path)
    assert rows == original
    assert {path.name for _, path in saved} == {"Fig5_violin.png", "Fig5_violin.pdf", "Fig5_violin.svg"}
    fig = saved[0][0]
    expected_order = ["BrainGNN", "BNT", "LG-GNN", "Com-brainTF", "BrainLM", "TFF", "SwiFT", "NeuroSTORM"]
    for ax, metric in zip(fig.axes, ["top1", "mAP"]):
        assert [label.get_text() for label in ax.get_yticklabels()] == expected_order
        for row in rows:
            model_id = row["model"].lower().replace("-", "")
            actual = []
            for collection in ax.collections:
                if collection.get_gid() in (f"scans-inliers-{metric}-{model_id}",
                                            f"scans-outliers-{metric}-{model_id}"):
                    actual.extend(collection.get_offsets()[:, 0])
            np.testing.assert_array_equal(sorted(actual), sorted(point[metric + "_percent"] for point in row["points"]))


@pytest.fixture
def source_workbook(tmp_path):
    path = tmp_path / "source_data.xlsx"
    book = openpyxl.Workbook()
    book.remove(book.active)
    expected = []
    for sheet_name, offset in (("Top1", 0), ("map", 2)):
        sheet = book.create_sheet(sheet_name)
        sheet.append(["HCPYA", *plot_fig5.MODELS])
        values = np.array([[80 + offset + row + col / 10 for col in range(8)]
                           for row in range(5)])
        for trial, numbers in enumerate(values, start=1):
            sheet.append([f"diameter {trial}", *numbers])
        expected.append(values)
    book.save(path)
    book.close()
    return path, expected


def test_source_data_cells_are_the_plotted_trial_values(source_workbook, tmp_path, monkeypatch):
    path, expected = source_workbook
    before = path.read_bytes()
    panels = plot_fig5.load_source_data(path)
    saved = []
    monkeypatch.setattr(Figure, "savefig", lambda fig, path, **kwargs: saved.append((fig, path)))
    plot_fig5.draw(panels, tmp_path, mode="points")
    assert path.read_bytes() == before
    assert {path.name for _, path in saved} == {"Fig5_points.png", "Fig5_points.pdf", "Fig5_points.svg"}
    fig = saved[0][0]
    for ax, values in zip(fig.axes, expected):
        assert ax.get_ylim() == (75., 100.)
        assert [label.get_text() for label in ax.get_xticklabels()] == [f"Trial {i}" for i in range(1, 6)]
        for index, line in enumerate(ax.lines):
            assert line.get_label() == plot_fig5.MODELS[index]
            np.testing.assert_array_equal(line.get_xdata(), np.arange(1, 6))
            np.testing.assert_array_equal(line.get_ydata(), values[:, index])


def test_invalid_source_data_score_is_rejected(source_workbook):
    path, _ = source_workbook
    book = openpyxl.load_workbook(path)
    book["Top1"]["B2"] = 101
    book.save(path)
    book.close()
    with pytest.raises(ValueError, match="Top1!B2"):
        plot_fig5.load_source_data(path)


def test_points_cli_reads_excel_without_ranking_lists(source_workbook, tmp_path, monkeypatch, capsys):
    path, expected = source_workbook
    drawn = []
    monkeypatch.setattr(sys, "argv", ["plot_fig5.py", "--mode", "points", "--source", str(path),
                                     "--output-dir", str(tmp_path / "output")])
    monkeypatch.setattr(plot_fig5, "register_paper_fonts", lambda mode: None)
    monkeypatch.setattr(plot_fig5, "draw", lambda data, output, mode: drawn.append((data, mode)))
    plot_fig5.main()
    assert len(drawn) == 1 and drawn[0][1] == "points"
    np.testing.assert_array_equal(drawn[0][0][0]["values"], expected[0])
    assert "Source Data mean" in capsys.readouterr().out


@pytest.mark.parametrize("mode,required", [("points", "--source"), ("violin", "--ranking-dir")])
def test_each_mode_requires_its_own_input(mode, required, monkeypatch, capsys):
    monkeypatch.setattr(sys, "argv", ["plot_fig5.py", "--mode", mode])
    with pytest.raises(SystemExit) as exc:
        plot_fig5.main()
    assert exc.value.code == 2
    assert f"{required} is required" in capsys.readouterr().err


@pytest.mark.parametrize("removed_args,error", [
    (["--mode", "query"], "invalid choice"),
    (["--mode", "simulated"], "invalid choice"),
    (["--mode", "both"], "invalid choice"),
    (["--simulation-seed", "42"], "unrecognized arguments"),
    (["--simulation-points", "400"], "unrecognized arguments"),
    (["--axis-range", "paper"], "unrecognized arguments"),
])
def test_removed_simulation_options_are_rejected(removed_args, error, tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(sys, "argv", ["plot_fig5.py", "--ranking-dir", str(tmp_path), *removed_args])
    with pytest.raises(SystemExit) as exc:
        plot_fig5.main()
    assert exc.value.code == 2
    assert error in capsys.readouterr().err
